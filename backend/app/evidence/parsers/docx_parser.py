"""Parser for Microsoft Word (.docx) files."""

import io
from typing import Tuple
from app.core.exceptions import MalformedFileError
from app.utils.text import clean_text


class DocxParser:
    """Extracts text paragraphs and tables from DOCX files."""

    def parse(self, content_bytes: bytes, filename: str) -> Tuple[str, dict]:
        try:
            import docx
            doc = docx.Document(io.BytesIO(content_bytes))
            parts = []

            # Paragraphs
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())

            # Tables
            table_count = len(doc.tables)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        parts.append(" | ".join(row_cells))

            full_text = clean_text("\n\n".join(parts))
            metadata = {
                "format": "docx",
                "paragraph_count": len(doc.paragraphs),
                "table_count": table_count,
                "character_count": len(full_text),
            }
            return full_text, metadata
        except Exception as e:
            raise MalformedFileError(filename, f"Failed to extract DOCX content: {str(e)}")
