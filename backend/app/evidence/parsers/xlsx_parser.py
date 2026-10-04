"""Parser for Excel spreadsheet (.xlsx) files."""

import io
from typing import Tuple
from app.core.exceptions import MalformedFileError
from app.utils.text import clean_text


class XLSXParser:
    """Extracts worksheets, tables, and rows from XLSX workbooks."""

    def parse(self, content_bytes: bytes, filename: str) -> Tuple[str, dict]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True, read_only=True)
            sheet_lines = []
            total_rows = 0

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet_lines.append(f"--- Sheet: {sheet_name} ---")
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue

                headers = [str(cell).strip() if cell is not None else f"col_{i}" for i, cell in enumerate(rows[0])]
                sheet_lines.append("Columns: " + ", ".join(headers))

                for row in rows[1:]:
                    if not any(row):
                        continue
                    total_rows += 1
                    cells = [f"{headers[i] if i < len(headers) else f'col_{i}'}: {str(c).strip()}" for i, c in enumerate(row) if c is not None]
                    sheet_lines.append(" | ".join(cells))

            wb.close()
            full_text = clean_text("\n".join(sheet_lines))
            metadata = {
                "format": "xlsx",
                "sheet_names": wb.sheetnames,
                "row_count": total_rows,
                "character_count": len(full_text),
            }
            return full_text, metadata
        except Exception as e:
            raise MalformedFileError(filename, f"Failed to parse XLSX workbook: {str(e)}")
